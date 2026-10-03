"""
excel_to_calc.py: turn Excel sheets into Obsidian ```calc blocks
=================================================================

Each sheet becomes one calc block. Formulas are kept, so the block
calculates live in Obsidian (Lab Calc plugin). References to other
sheets (Sheet2!B3) become references to the matching block (sheet2!B3).

Usage (Windows: use "py" instead of "python" if that's how Python is set up)
-----
    python excel_to_calc.py "My calcs.xlsx"                  all sheets
    python excel_to_calc.py "My calcs.xlsx" --sheet Recipe   one sheet
    python excel_to_calc.py "My calcs.xlsx" --sheet Recipe --range A1:G8
    python excel_to_calc.py "My calcs.xlsx" --name sol1      set block name (one sheet only)

Output
------
- Copied to the clipboard, ready to paste into a note.
- Also saved next to the Excel file as  <file>.calc.md
- Warnings (unsupported functions, etc.) are printed.

Needs:  pip install openpyxl
Google Sheets: File > Download > Microsoft Excel (.xlsx) first.
"""

import argparse
import datetime as _dt
import platform
import re
import subprocess
import sys
from pathlib import Path

try:
    from openpyxl import load_workbook
    from openpyxl.utils import column_index_from_string, get_column_letter, range_boundaries
except ImportError:
    sys.exit("openpyxl is missing. Install it with:  pip install openpyxl")

# Functions the Lab Calc plugin understands. Keep in sync with main.js.
SUPPORTED = {
    "SUM", "PRODUCT", "AVERAGE", "MIN", "MAX", "COUNT", "COUNTA", "SUMPRODUCT",
    "ROUND", "ROUNDUP", "ROUNDDOWN", "INT", "ABS", "SQRT", "EXP", "LN", "LOG10",
    "LOG", "POWER", "MOD", "PI", "IF", "IFERROR", "AND", "OR", "NOT", "AVG",
    "XLOOKUP", "MATCH", "INDEX", "CONCAT",
}

SHEET_REF = re.compile(r"(?:'((?:[^']|'')+)'|([A-Za-z_][\w.]*))!")
CELL_REF = re.compile(r"(?<![A-Za-z_.!])(\$?)([A-Z]{1,3})(\$?)(\d+)(?![\w(])")
FUNC = re.compile(r"([A-Z][A-Z0-9.]*)\s*\(")


def slug(name: str) -> str:
    s = re.sub(r"[^0-9A-Za-z_]+", "_", name.strip()).strip("_").lower()
    if not s or not s[0].isalpha():
        s = "s_" + s
    return s


def fmt_value(v) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, float):
        return repr(int(v)) if v.is_integer() else repr(v)
    if isinstance(v, int):
        return str(v)
    if isinstance(v, (_dt.datetime, _dt.date)):
        return v.isoformat()[:10] if isinstance(v, _dt.date) and not isinstance(v, _dt.datetime) else v.isoformat(sep=" ", timespec="minutes")
    return str(v).replace("\n", " ").strip()


def split_strings(formula: str):
    """Yield (is_string_literal, text) chunks so refs inside "quotes" are left alone."""
    out, buf, in_str, i = [], "", False, 0
    while i < len(formula):
        ch = formula[i]
        if ch == '"':
            if in_str and i + 1 < len(formula) and formula[i + 1] == '"':
                buf += '""'; i += 2; continue
            if not in_str:
                out.append((False, buf)); buf = '"'; in_str = True
            else:
                buf += '"'; out.append((True, buf)); buf = ""; in_str = False
            i += 1; continue
        buf += ch; i += 1
    out.append((in_str, buf))
    return out


def convert_formula(f: str, this_sheet: str, slugs: dict, offsets: dict, warnings: list, where: str) -> str:
    """Rewrite an Excel formula for Lab Calc."""
    def shift(sheet_key, col_letters, row_digits):
        dc, dr = offsets.get(sheet_key, (0, 0))
        c = column_index_from_string(col_letters) - dc
        r = int(row_digits) - dr
        if c < 1 or r < 1:
            warnings.append(f"{where}: reference {col_letters}{row_digits} is outside the exported range")
            return f"{col_letters}{row_digits}"
        return f"{get_column_letter(c)}{r}"

    pieces = []
    for is_str, chunk in split_strings(f):
        if is_str:
            pieces.append(chunk); continue
        if "[" in chunk:
            warnings.append(f"{where}: links to another workbook are not supported ({f})")
        # 1) sheet-qualified refs:  Sheet2!A1 / 'My sheet'!A1:B4
        res, pos = "", 0
        for m in SHEET_REF.finditer(chunk):
            res += chunk[pos:m.start()]
            sheet = (m.group(1) or m.group(2)).replace("''", "'")
            if sheet not in slugs:
                warnings.append(f"{where}: refers to sheet '{sheet}' that isn't being exported")
            key = sheet
            res += slugs.get(sheet, slug(sheet)) + "!"
            # rewrite the ref(s) right after the qualifier with that sheet's offset
            tail = chunk[m.end():]
            rm = re.match(r"\$?([A-Z]{1,3})\$?(\d+)(?::\$?([A-Z]{1,3})\$?(\d+))?", tail)
            if rm:
                a = shift(key, rm.group(1), rm.group(2))
                b = (":" + shift(key, rm.group(3), rm.group(4))) if rm.group(3) else ""
                res += "\x00" + a + b + "\x00"
                pos = m.end() + rm.end()
            else:
                pos = m.end()
        res += chunk[pos:]
        # 2) plain refs on this sheet (skip those already handled, marked with \x00)
        def plain(mm):
            if offsets.get(this_sheet, (0, 0)) == (0, 0):
                return mm.group(0)
            new = re.match(r"([A-Z]+)(\d+)", shift(this_sheet, mm.group(2), mm.group(4)))
            return mm.group(1) + new.group(1) + mm.group(3) + new.group(2)

        parts = res.split("\x00")
        for i in range(0, len(parts), 2):
            parts[i] = CELL_REF.sub(plain, parts[i])
        pieces.append("".join(parts))
    out = "".join(pieces)

    for fn in FUNC.findall(out.upper()):
        name = fn.replace("_XLFN.", "")
        if name not in SUPPORTED:
            warnings.append(f"{where}: function {name}() isn't supported by Lab Calc yet")
    return out.replace("_xlfn.", "")


def sheet_bounds(ws, rng):
    if rng:
        min_col, min_row, max_col, max_row = range_boundaries(rng.upper())
        return min_row, min_col, max_row, max_col
    max_row = max_col = 0
    for row in ws.iter_rows():
        for cell in row:
            if cell.value not in (None, ""):
                max_row = max(max_row, cell.row)
                max_col = max(max_col, cell.column)
    return 1, 1, max_row, max_col     # always start at A1 so references don't move


def build_block(ws, name, bounds, slugs, offsets, warnings):
    r0, c0, r1, c1 = bounds
    lines = [f"name: {name}", f"title: {ws.title}"]
    rows = []
    for r in range(r0, r1 + 1):
        row = []
        for c in range(c0, c1 + 1):
            v = ws.cell(row=r, column=c).value
            where = f"{ws.title}!{get_column_letter(c)}{r}"
            if hasattr(v, "text"):                       # ArrayFormula
                v = v.text
            if isinstance(v, str) and v.startswith("="):
                v = convert_formula(v, ws.title, slugs, offsets, warnings, where)
            row.append(fmt_value(v).replace("|", "\\|"))
        rows.append(row)
    if not rows:
        return None
    widths = [max(3, *(len(rw[i]) for rw in rows)) for i in range(len(rows[0]))]
    def line(rw):
        return "| " + " | ".join(cell.ljust(widths[i]) for i, cell in enumerate(rw)) + " |"
    lines.append(line(rows[0]))
    lines.append("| " + " | ".join("-" * w for w in widths) + " |")
    lines += [line(rw) for rw in rows[1:]]
    return "```calc\n" + "\n".join(lines) + "\n```"


def copy_to_clipboard(text: str) -> bool:
    try:
        system = platform.system()
        if system == "Windows":
            subprocess.run(["clip"], input=text.encode("utf-16le"), check=True)
        elif system == "Darwin":
            subprocess.run(["pbcopy"], input=text.encode("utf-8"), check=True)
        else:
            subprocess.run(["xclip", "-selection", "clipboard"], input=text.encode("utf-8"), check=True)
        return True
    except Exception:
        return False


def main():
    ap = argparse.ArgumentParser(description="Convert Excel sheets to Obsidian ```calc blocks")
    ap.add_argument("xlsx", help="path to the .xlsx file")
    ap.add_argument("--sheet", action="append", help="sheet to export (repeat for several). Default: all")
    ap.add_argument("--range", help="cell range to export, e.g. A1:G8 (one sheet only)")
    ap.add_argument("--name", help="block name to use (one sheet only)")
    ap.add_argument("--no-clipboard", action="store_true")
    args = ap.parse_args()

    path = Path(args.xlsx)
    if not path.exists():
        sys.exit(f"File not found: {path}")
    wb = load_workbook(path, data_only=False)

    sheets = args.sheet or wb.sheetnames
    missing = [s for s in sheets if s not in wb.sheetnames]
    if missing:
        sys.exit(f"Sheet(s) not found: {missing}. Available: {wb.sheetnames}")
    if (args.range or args.name) and len(sheets) != 1:
        sys.exit("--range and --name need exactly one --sheet")

    slugs = {s: slug(s) for s in wb.sheetnames}
    if args.name:
        slugs[sheets[0]] = slug(args.name)

    bounds, offsets = {}, {}
    for s in sheets:
        b = sheet_bounds(wb[s], args.range)
        bounds[s] = b
        offsets[s] = (b[1] - 1, b[0] - 1)       # (cols, rows) shifted by --range

    warnings, blocks = [], []
    for s in sheets:
        blk = build_block(wb[s], slugs[s], bounds[s], slugs, offsets, warnings)
        if blk:
            blocks.append(blk)
        else:
            warnings.append(f"Sheet '{s}' is empty, skipped")

    text = "\n\n".join(blocks) + "\n"
    out = path.with_suffix(".calc.md")
    out.write_text(text, encoding="utf-8")

    print(text)
    print(f"Saved to: {out}")
    if not args.no_clipboard and copy_to_clipboard(text):
        print("Copied to clipboard ✔  Paste it into your note.")
    if warnings:
        print("\nWarnings:")
        for w in dict.fromkeys(warnings):
            print("  -", w)
    print("\nTip: Excel shows percentages as 15% but stores 0.15; check any % cells.")


if __name__ == "__main__":
    main()
